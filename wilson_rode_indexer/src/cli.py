"""Command-line interface and phase orchestration.

Commands
--------
``locate``    Search likely Google Drive locations for the source folder.
``doctor``    Report dependency and configuration health.
``phase1``    Structural inventory (no substantive text is read).
``phase2``    Text and page-condition audit; pilot sample by default.
``phase3``    Derived metadata, duplicates, families, tags, and workbooks.
``qc``        Emit the quality-control sample for manual inspection.
``search``    Query the FTS5 index (prints file identities, never text).
``status``    Show processing progress.
``rerun-failed`` Requeue only the files that errored.

The source folder is opened read-only by every command. Nothing beneath it
is created, renamed, moved, modified, or deleted.
"""

from __future__ import annotations

import argparse
import os
import random
import sqlite3
import sys
import uuid
from collections import Counter
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from src import bates as bates_mod
from src import document_family, duplicates, issue_tags, metadata_extract, pdf_extract, reports
from src.config import Config, ConfigError, load_config
from src.database import STATUS_DONE, STATUS_ERROR, Database, utcnow
from src.inventory import run_inventory
from src.logging_setup import get_logger, setup_logging
from src.version import APP_NAME, APP_VERSION

LOG = get_logger("cli")

#: Directories searched by ``wri locate``, in order.
MAC_DRIVE_ROOTS = (
    "~/Library/CloudStorage",
    "~/Google Drive",
    "~/Google Drive File Stream",
    "/Volumes/GoogleDrive",
    "~/Documents",
    "~/Desktop",
)

TARGET_FOLDER_NAME = "Wilson-Rode File"
DERIVED_FOLDER_NAME = "Wilson_Rode_Derived_Index"


# ---------------------------------------------------------------------------
# locate
# ---------------------------------------------------------------------------


def find_source_folders(
    target: str = TARGET_FOLDER_NAME, *, roots: Sequence[str] = MAC_DRIVE_ROOTS, max_depth: int = 6
) -> list[Path]:
    """Search likely Google Drive locations for the production folder.

    Args:
        target: Exact directory name to find.
        roots: Candidate search roots (``~`` is expanded).
        max_depth: Levels to descend below each root.

    Returns:
        Matching directories, deduplicated, in discovery order.
    """
    found: list[Path] = []
    seen: set[Path] = set()
    for raw_root in roots:
        root = Path(os.path.expanduser(raw_root))
        if not root.is_dir():
            continue
        base_depth = len(root.parts)
        for dirpath, dirnames, _filenames in os.walk(root, followlinks=False):
            current = Path(dirpath)
            if len(current.parts) - base_depth >= max_depth:
                dirnames[:] = []
                continue
            for name in list(dirnames):
                if name == target:
                    candidate = (current / name).resolve()
                    if candidate not in seen:
                        seen.add(candidate)
                        found.append(candidate)
    return found


def command_locate(args: argparse.Namespace) -> int:
    """Locate the source folder and print candidate paths."""
    target = args.name
    print(f"Searching for a folder named exactly: {target}")
    for root in MAC_DRIVE_ROOTS:
        expanded = Path(os.path.expanduser(root))
        print(f"  {'[present]' if expanded.is_dir() else '[missing]'} {expanded}")

    matches = find_source_folders(target)
    if not matches:
        print("\nNo matching folder found.")
        print("If the production lives elsewhere, set paths.source_folder in config.yaml")
        print("to its absolute path, then run: wri doctor")
        return 1

    print(f"\nFound {len(matches)} matching folder(s):\n")
    for index, path in enumerate(matches, start=1):
        try:
            entries = sum(1 for _ in path.iterdir())
        except OSError:
            entries = -1
        print(f"  [{index}] {path}")
        print(f"      top-level entries: {entries if entries >= 0 else 'unreadable'}")
        print(f"      suggested output:  {path.parent / DERIVED_FOLDER_NAME}")

    if len(matches) > 1:
        print(
            "\nMultiple folders match. Choose one and set paths.source_folder and "
            "paths.output_folder in config.yaml accordingly. Nothing has been modified."
        )
        return 2

    print("\nTo use this folder, set in config.yaml:")
    print(f"  paths.source_folder: \"{matches[0]}\"")
    print(f"  paths.output_folder: \"{matches[0].parent / DERIVED_FOLDER_NAME}\"")
    return 0


# ---------------------------------------------------------------------------
# doctor
# ---------------------------------------------------------------------------


def command_doctor(args: argparse.Namespace) -> int:
    """Report dependency, configuration, and permission health."""
    problems = 0
    print(f"{APP_NAME} v{APP_VERSION}")
    print(f"Python {sys.version.split()[0]} at {sys.executable}\n")

    print("Required libraries:")
    for module, label in (
        ("fitz", "PyMuPDF"),
        ("pypdf", "pypdf"),
        ("pandas", "pandas"),
        ("openpyxl", "openpyxl"),
        ("dateutil", "python-dateutil"),
        ("rapidfuzz", "rapidfuzz"),
        ("tqdm", "tqdm"),
        ("PIL", "Pillow"),
        ("yaml", "PyYAML"),
    ):
        try:
            __import__(module)
            print(f"  [ok]      {label}")
        except ImportError:
            problems += 1
            print(f"  [MISSING] {label}")

    available, message = pdf_extract.ocr_available()
    print(f"\nOCR: {'[ok]' if available else '[optional, not installed]'} {message}")
    if not available:
        print("  Install locally with: brew install ocrmypdf tesseract")

    try:
        connection = sqlite3.connect(":memory:")
        connection.execute("CREATE VIRTUAL TABLE t USING fts5(a)")
        connection.close()
        print("SQLite FTS5: [ok]")
    except sqlite3.Error as exc:
        problems += 1
        print(f"SQLite FTS5: [MISSING] {exc}")

    print("\nConfiguration:")
    try:
        config = load_config(args.config)
        print(f"  config file:   {config.config_path}")
        print(f"  source folder: {config.paths.source_folder}")
        print(f"  output folder: {config.paths.output_folder}")
        print(f"  workers:       {config.worker_count()}")
        config.validate_paths(require_source=True)
        print("  [ok] Source exists and the output folder is safely outside it.")

        readable = os.access(config.paths.source_folder, os.R_OK)
        print(f"  source readable: {'[ok]' if readable else '[NO READ ACCESS]'}")
        if not readable:
            problems += 1
    except ConfigError as exc:
        problems += 1
        print(f"  [PROBLEM] {exc}")

    print(f"\n{'All checks passed.' if problems == 0 else f'{problems} problem(s) found.'}")
    return 0 if problems == 0 else 1


# ---------------------------------------------------------------------------
# phase 1
# ---------------------------------------------------------------------------


def command_phase1(args: argparse.Namespace) -> int:
    """Run the structural inventory and write the Phase 1 deliverables."""
    config = _prepare(args)
    log_path = setup_logging(
        config.paths.logs_dir,
        run_name="phase1",
        console_level=str(config.get("logging.console_level", "INFO")),
        file_level=str(config.get("logging.file_level", "DEBUG")),
        max_bytes=int(config.get("logging.max_bytes", 20971520)),
        backup_count=int(config.get("logging.backup_count", 10)),
    )

    with Database(config.paths.db_path) as database:
        stats = run_inventory(
            config,
            database,
            force=args.force,
            limit=args.limit,
            progress=not args.no_progress,
        )
        written = reports.write_phase1_reports(config, database, stats.as_dict())

    _print_phase1_summary(stats.as_dict(), written, log_path, config)
    return 0


def _print_phase1_summary(
    stats: Mapping[str, Any], written: Mapping[str, Path], log_path: Path, config: Config
) -> None:
    """Print aggregate Phase 1 results. Contains no document content."""
    total_bytes = int(stats.get("total_bytes", 0) or 0)
    print("\n" + "=" * 72)
    print("PHASE 1 COMPLETE - STRUCTURAL INVENTORY ONLY")
    print("=" * 72)
    print(f"Source (read-only): {config.paths.source_folder}")
    print(f"Derived index:      {config.paths.output_folder}")
    print()
    print(f"  Total files ................. {stats.get('total_files', 0):,}")
    print(f"  PDFs ........................ {stats.get('total_pdfs', 0):,}")
    print(f"  Non-PDF files ............... {stats.get('total_non_pdfs', 0):,}")
    print(f"  Total storage ............... {total_bytes / (1024 ** 3):.2f} GiB")
    print(f"  Total pages ................. {stats.get('total_pages', 0):,}")
    print(f"  Processed this run .......... {stats.get('processed_now', 0):,}")
    print(f"  Skipped (already complete) .. {stats.get('skipped_resumed', 0):,}")
    print()
    print(f"  Min apparent Bates .......... {stats.get('min_bates')}")
    print(f"  Max apparent Bates .......... {stats.get('max_bates')}")
    print(f"  PDFs without a Bates name ... {stats.get('bates_missing', 0):,}")
    print(f"  Non-compliant filenames ..... {stats.get('noncompliant_filenames', 0):,}")
    print(f"  Preliminary gap runs ........ {stats.get('filename_gap_runs', 0):,}")
    print(f"  Preliminary missing numbers . {stats.get('filename_gap_missing_numbers', 0):,}")
    print("    (preliminary: multi-page PDFs consume intervening Bates numbers)")
    print()
    print(f"  Duplicate filenames ......... {stats.get('duplicate_filename_files', 0):,}"
          f" in {stats.get('duplicate_filename_groups', 0):,} group(s)")
    print(f"  Exact duplicate files ....... {stats.get('exact_duplicate_files', 0):,}"
          f" in {stats.get('exact_duplicate_groups', 0):,} group(s)")
    print(f"  Encrypted PDFs .............. {stats.get('encrypted', 0):,}")
    print(f"  Corrupt / unreadable ........ {stats.get('corrupt', 0):,}")
    print(f"  Cloud placeholders .......... {stats.get('cloud_placeholders', 0):,}")
    print(f"  Processing errors ........... {stats.get('errors', 0):,}")

    distribution = stats.get("page_count_distribution", {}) or {}
    if distribution:
        print("\n  Page-count distribution:")
        for bucket, count in distribution.items():
            print(f"    {bucket:>8} pages : {count:,}")

    print("\n  Reports written:")
    for name, path in written.items():
        print(f"    {name}")
    print(f"\n  Log: {log_path}")
    print("\nSTOPPED AFTER PHASE 1 as instructed. Phase 2 requires explicit approval.")
    print("=" * 72)


# ---------------------------------------------------------------------------
# phase 2
# ---------------------------------------------------------------------------


def select_pilot_sample(database: Database, config: Config, *, size: int = 200) -> list[int]:
    """Choose a stratified pilot sample of documents.

    The sample deliberately over-represents difficult cases: low and high
    Bates ranges, very small and very large files, single- and multi-page
    documents, likely image-only and likely blank PDFs, exact duplicates,
    and malformed filenames. A uniform random sample would under-sample
    exactly the cases most likely to break extraction.

    Args:
        database: Populated processing-state database.
        config: Loaded configuration (for the deterministic seed).
        size: Target sample size.

    Returns:
        Distinct file ids, capped at ``size``.
    """
    rng = random.Random(int(config.get("qc.random_seed", 20240101)))
    strata: list[tuple[str, str, int]] = [
        ("lowest Bates", "is_pdf=1 AND bates_filename_num IS NOT NULL ORDER BY bates_filename_num ASC", 20),
        ("highest Bates", "is_pdf=1 AND bates_filename_num IS NOT NULL ORDER BY bates_filename_num DESC", 20),
        ("smallest files", "is_pdf=1 AND file_size IS NOT NULL AND is_corrupt=0 ORDER BY file_size ASC", 20),
        ("largest files", "is_pdf=1 AND file_size IS NOT NULL ORDER BY file_size DESC", 20),
        ("single page", "is_pdf=1 AND page_count = 1 ORDER BY RANDOM()", 25),
        ("multi page", "is_pdf=1 AND page_count > 1 ORDER BY RANDOM()", 25),
        ("many pages", "is_pdf=1 AND page_count >= 10 ORDER BY page_count DESC", 15),
        ("likely image-only or blank", "is_pdf=1 AND page_count >= 1 AND file_size < 60000 ORDER BY RANDOM()", 20),
        ("exact duplicates", "is_pdf=1 AND exact_dup_group IS NOT NULL ORDER BY exact_dup_group", 15),
        ("malformed names", "is_pdf=1 AND (filename_compliant = 0 OR bates_filename_num IS NULL) ORDER BY RANDOM()", 15),
        ("unreadable or encrypted", "is_pdf=1 AND (is_corrupt=1 OR is_encrypted=1) ORDER BY RANDOM()", 5),
    ]

    selected: list[int] = []
    seen: set[int] = set()
    for label, predicate, quota in strata:
        rows = database.query(f"SELECT id FROM files WHERE {predicate} LIMIT ?", (quota * 3,))
        picked = 0
        for row in rows:
            file_id = int(row["id"])
            if file_id in seen:
                continue
            seen.add(file_id)
            selected.append(file_id)
            picked += 1
            if picked >= quota:
                break
        LOG.info("Pilot stratum selected", extra={"stratum": label, "selected": picked})

    if len(selected) < size:
        filler = database.query(
            "SELECT id FROM files WHERE is_pdf=1 ORDER BY RANDOM() LIMIT ?", (size * 2,)
        )
        for row in filler:
            file_id = int(row["id"])
            if file_id not in seen:
                seen.add(file_id)
                selected.append(file_id)
            if len(selected) >= size:
                break

    rng.shuffle(selected)
    return selected[:size]


def _extract_documents(
    config: Config,
    database: Database,
    rows: Sequence[sqlite3.Row],
    *,
    run_id: str,
    progress: bool = True,
    force: bool = False,
) -> dict[str, Any]:
    """Extract text and page conditions for a set of files.

    Returns aggregate statistics. Extracted text goes to FTS5 only.
    """
    from tqdm import tqdm

    thresholds = pdf_extract.thresholds_from_config(config)
    worker_kwargs = {
        "thresholds": thresholds,
        "bates_page_pattern": str(config.get("bates.page_regex")),
        "bates_prefix": str(config.get("bates.prefix", "WILSONRODE")),
        "bates_digits": int(config.get("bates.expected_digits", 6)),
        "max_pages": int(config.get("performance.max_pages_per_document", 0) or 0),
        "max_indexed_chars": int(
            config.get("extraction.max_indexed_chars_per_document", 0) or 0
        ),
    }

    todo: list[sqlite3.Row] = []
    skipped = 0
    for row in rows:
        if force:
            todo.append(row)
            continue
        if database.needs_processing(
            str(row["rel_path"]),
            size=int(row["file_size"] or 0),
            mtime_iso=str(row["modified_ts"] or ""),
            stage="extraction_status",
        ):
            todo.append(row)
        else:
            skipped += 1

    stats: dict[str, Any] = {
        "run_id": run_id,
        "documents": len(rows),
        "documents_processed": len(todo),
        "documents_skipped_resumed": skipped,
        "pages": 0,
        "failed": 0,
        "page_conditions": Counter(),
        "blank_subtypes": Counter(),
        "pages_with_observed_bates": 0,
        "bates_filename_agreement": 0,
        "bates_filename_disagreement": 0,
        "ocr_required_pages": 0,
        "ocr_completed_pages": 0,
        "ocr_failures": 0,
    }

    workers = config.worker_count()
    bar = tqdm(total=len(todo), disable=not progress, unit="doc", desc="Extracting")

    def absorb(row: sqlite3.Row, payload: dict[str, Any]) -> None:
        """Persist one extraction result and fold it into the statistics."""
        rel_path = str(row["rel_path"])
        pdf_extract.persist_extraction(
            database, int(row["id"]), rel_path, payload, run_id=run_id
        )
        if payload.get("status") == STATUS_ERROR:
            stats["failed"] += 1
            LOG.warning(
                "Extraction failed",
                extra={"rel_path": rel_path, "error": payload.get("error_message")},
            )
            return
        pages = payload.get("pages") or []
        stats["pages"] += len(pages)
        for page in pages:
            stats["page_conditions"][page["condition"]] += 1
            if page.get("blank_subtype"):
                stats["blank_subtypes"][page["blank_subtype"]] += 1
            if page.get("bates_observed_num") is not None:
                stats["pages_with_observed_bates"] += 1
        stats["ocr_required_pages"] += int(payload.get("ocr_required_pages", 0))

        observed = payload.get("observed_bates") or []
        filename_number = row["bates_filename_num"]
        if observed and filename_number is not None:
            if int(filename_number) in {int(value) for value in observed}:
                stats["bates_filename_agreement"] += 1
            else:
                stats["bates_filename_disagreement"] += 1
                LOG.warning(
                    "Observed Bates disagrees with the filename",
                    extra={
                        "rel_path": rel_path,
                        "filename_bates": int(filename_number),
                        "observed_first": int(observed[0]),
                    },
                )

    try:
        if workers <= 1 or len(todo) <= 1:
            for row in todo:
                absorb(row, pdf_extract.extract_document(
                    str(row["abs_path"]), str(row["rel_path"]), **worker_kwargs
                ))
                bar.update(1)
        else:
            with ProcessPoolExecutor(max_workers=workers) as pool:
                futures = {
                    pool.submit(
                        pdf_extract.extract_document,
                        str(row["abs_path"]),
                        str(row["rel_path"]),
                        **worker_kwargs,
                    ): row
                    for row in todo
                }
                for future in as_completed(futures):
                    row = futures[future]
                    try:
                        absorb(row, future.result())
                    except Exception as exc:  # noqa: BLE001
                        stats["failed"] += 1
                        LOG.error(
                            "Extraction worker raised",
                            extra={"rel_path": str(row["rel_path"]), "error": str(exc)},
                        )
                        database.update_file(
                            str(row["rel_path"]),
                            {
                                "extraction_status": STATUS_ERROR,
                                "error_message": f"{type(exc).__name__}: {exc}"[:500],
                                "last_processed_ts": utcnow(),
                            },
                        )
                    bar.update(1)
    finally:
        bar.close()

    stats["page_conditions"] = dict(stats["page_conditions"])
    stats["blank_subtypes"] = dict(stats["blank_subtypes"])
    return stats


def _run_ocr_pass(
    config: Config, database: Database, file_ids: Sequence[int], *, run_id: str
) -> dict[str, Any]:
    """OCR the documents that need it, writing only to the cache folder."""
    ocr_config = config.section("ocr")
    result = {"attempted": 0, "completed": 0, "failed": 0, "skipped": 0}

    if not ocr_config.get("enabled", True) or not ocr_config.get("run_ocr", True):
        LOG.info("OCR disabled in configuration; candidates reported only")
        result["skipped"] = len(file_ids)
        return result

    available, message = pdf_extract.ocr_available()
    if not available:
        LOG.warning("OCR unavailable; candidates reported only", extra={"detail": message})
        result["skipped"] = len(file_ids)
        return result

    max_bytes = int(ocr_config.get("max_source_mib", 200)) * 1024 * 1024
    marks = ",".join("?" for _ in file_ids) or "NULL"
    candidates = database.query(
        "SELECT DISTINCT f.id, f.rel_path, f.abs_path, f.sha256, f.file_size"
        " FROM files f JOIN pages p ON p.file_id = f.id"
        f" WHERE p.ocr_status = 'required' AND f.id IN ({marks})",
        list(file_ids),
    )

    for row in candidates:
        result["attempted"] += 1
        if int(row["file_size"] or 0) > max_bytes:
            result["skipped"] += 1
            LOG.warning("Skipping OCR: source exceeds the size cap",
                        extra={"rel_path": str(row["rel_path"])})
            continue
        target = pdf_extract.ocr_cache_path(
            config.paths.ocr_cache_dir, str(row["rel_path"]), row["sha256"]
        )
        success, message = pdf_extract.run_ocr(
            Path(str(row["abs_path"])),
            target,
            language=str(ocr_config.get("language", "eng")),
            extra_args=list(ocr_config.get("ocrmypdf_args", ["--skip-text", "--quiet"])),
            timeout_seconds=int(ocr_config.get("timeout_seconds", 600)),
        )
        if success:
            result["completed"] += 1
            database.update_file(
                str(row["rel_path"]), {"ocr_status": pdf_extract.OcrStatus.COMPLETED.value}
            )
        else:
            result["failed"] += 1
            database.update_file(
                str(row["rel_path"]), {"ocr_status": pdf_extract.OcrStatus.FAILED.value}
            )
            LOG.warning("OCR failed", extra={"rel_path": str(row["rel_path"]), "detail": message})
        database.audit(
            "ocr", run_id=run_id, phase="phase2", rel_path=str(row["rel_path"]),
            detail={"success": success, "message": message, "cache": str(target)},
        )
    return result


def command_phase2(args: argparse.Namespace) -> int:
    """Run the text and page-condition audit."""
    config = _prepare(args)
    log_path = setup_logging(
        config.paths.logs_dir,
        run_name="phase2",
        console_level=str(config.get("logging.console_level", "INFO")),
        file_level=str(config.get("logging.file_level", "DEBUG")),
    )
    run_id = f"phase2-{uuid.uuid4().hex[:12]}"

    with Database(config.paths.db_path) as database:
        if database.count_files() == 0:
            print("No inventory found. Run 'wri phase1' first.")
            return 1

        database.start_run(
            run_id,
            "phase2",
            {
                "config_path": str(config.config_path),
                "source_folder": str(config.paths.source_folder),
                "output_folder": str(config.paths.output_folder),
            },
        )

        if args.full:
            rows = database.query(
                "SELECT * FROM files WHERE is_pdf = 1 AND is_corrupt = 0"
                " AND is_cloud_placeholder = 0 ORDER BY bates_filename_num, rel_path"
            )
            file_ids = [int(row["id"]) for row in rows]
            scope_label = "full production"
        else:
            size = args.sample or int(config.get("qc.pilot_sample_size", 200))
            file_ids = select_pilot_sample(database, config, size=size)
            marks = ",".join("?" for _ in file_ids) or "NULL"
            rows = database.query(
                f"SELECT * FROM files WHERE id IN ({marks}) ORDER BY rel_path", file_ids
            )
            scope_label = f"pilot sample of {len(file_ids)} document(s)"

        LOG.info("Phase 2 scope selected", extra={"scope": scope_label, "documents": len(rows)})
        print(f"Phase 2 scope: {scope_label}")

        stats = _extract_documents(
            config, database, rows, run_id=run_id,
            progress=not args.no_progress, force=args.force,
        )

        ocr_result = _run_ocr_pass(config, database, file_ids, run_id=run_id)
        stats["ocr_completed_pages"] = ocr_result["completed"]
        stats["ocr_failures"] = ocr_result["failed"]
        available, message = pdf_extract.ocr_available()
        stats["ocr_available"] = message
        stats["accuracy_notes"] = _accuracy_notes(stats)

        written = reports.write_phase2_reports(
            config, database, stats, file_ids=file_ids, pilot=not args.full
        )
        database.finish_run(run_id, {k: v for k, v in stats.items() if not isinstance(v, Counter)})

    _print_phase2_summary(stats, written, log_path, pilot=not args.full)
    return 0


def _accuracy_notes(stats: Mapping[str, Any]) -> list[str]:
    """Derive plain-language accuracy warnings from Phase 2 statistics."""
    notes: list[str] = []
    pages = int(stats.get("pages", 0) or 0)
    conditions = stats.get("page_conditions", {}) or {}

    if stats.get("failed"):
        notes.append(
            f"{stats['failed']} document(s) failed extraction entirely; see "
            "08_Extraction_Failures.xlsx."
        )
    if stats.get("bates_filename_disagreement"):
        notes.append(
            f"{stats['bates_filename_disagreement']} document(s) had a Bates stamp on the "
            "page that did not match the filename. Where they disagree, the page value "
            "governs, but these documents warrant manual inspection."
        )
    uncertain = int(conditions.get("Uncertain", 0) or 0)
    if pages and uncertain / pages > 0.05:
        notes.append(
            f"{uncertain:,} of {pages:,} pages ({uncertain / pages:.1%}) could not be "
            "classified confidently. Consider tuning the extraction thresholds in "
            "config.yaml before the full run."
        )
    ocr_required = int(stats.get("ocr_required_pages", 0) or 0)
    if pages and ocr_required / pages > 0.25:
        notes.append(
            f"{ocr_required:,} of {pages:,} pages ({ocr_required / pages:.1%}) need OCR. "
            "A full run will be substantially slower, and fields derived from OCR text "
            "carry recognition error."
        )
    if stats.get("ocr_failures"):
        notes.append(f"{stats['ocr_failures']} OCR attempt(s) failed.")
    if not notes:
        notes.append(
            "No systematic accuracy problem was detected automatically. Automated "
            "checks cannot replace the manual QC sample - run 'wri qc' and inspect it."
        )
    return notes


def _print_phase2_summary(
    stats: Mapping[str, Any], written: Mapping[str, Path], log_path: Path, *, pilot: bool
) -> None:
    """Print aggregate Phase 2 results. Contains no document content."""
    print("\n" + "=" * 72)
    print(f"PHASE 2 COMPLETE - {'PILOT' if pilot else 'FULL PRODUCTION'}")
    print("=" * 72)
    print(f"  Documents in scope .......... {stats.get('documents', 0):,}")
    print(f"  Documents processed ......... {stats.get('documents_processed', 0):,}")
    print(f"  Skipped (already complete) .. {stats.get('documents_skipped_resumed', 0):,}")
    print(f"  Pages examined .............. {stats.get('pages', 0):,}")
    print(f"  Extraction failures ......... {stats.get('failed', 0):,}")
    print("\n  Page conditions:")
    for condition, count in sorted((stats.get("page_conditions") or {}).items()):
        print(f"    {condition:<22} {count:,}")
    if stats.get("blank_subtypes"):
        print("\n  Blank / Bates-only subtypes:")
        for subtype, count in sorted(stats["blank_subtypes"].items()):
            print(f"    {subtype:<32} {count:,}")
    print(f"\n  Pages with an observed Bates stamp ... {stats.get('pages_with_observed_bates', 0):,}")
    print(f"  Filename/page Bates agreement ....... {stats.get('bates_filename_agreement', 0):,}")
    print(f"  Filename/page Bates DISAGREEMENT .... {stats.get('bates_filename_disagreement', 0):,}")
    print(f"\n  OCR pages required .................. {stats.get('ocr_required_pages', 0):,}")
    print(f"  OCR completed ....................... {stats.get('ocr_completed_pages', 0):,}")
    print(f"  OCR failures ........................ {stats.get('ocr_failures', 0):,}")
    print("\n  Accuracy notes:")
    for note in stats.get("accuracy_notes", []):
        print(f"    - {note}")
    print("\n  Reports written:")
    for name in written:
        print(f"    {name}")
    print(f"\n  Log: {log_path}")
    if pilot:
        print("\nSTOPPED AFTER THE PILOT. Review the accuracy notes above before")
        print("approving a full-production run.")
    print("=" * 72)


# ---------------------------------------------------------------------------
# phase 3
# ---------------------------------------------------------------------------


def command_phase3(args: argparse.Namespace) -> int:
    """Derive metadata for the production and write workbooks 09-20."""
    config = _prepare(args)
    log_path = setup_logging(
        config.paths.logs_dir,
        run_name="phase3",
        console_level=str(config.get("logging.console_level", "INFO")),
        file_level=str(config.get("logging.file_level", "DEBUG")),
    )
    run_id = f"phase3-{uuid.uuid4().hex[:12]}"

    with Database(config.paths.db_path) as database:
        if database.count_files() == 0:
            print("No inventory found. Run 'wri phase1' first.")
            return 1

        database.start_run(
            run_id,
            "phase3",
            {
                "config_path": str(config.config_path),
                "source_folder": str(config.paths.source_folder),
                "output_folder": str(config.paths.output_folder),
            },
        )

        # Ensure text exists for everything in scope.
        rows = database.query(
            "SELECT * FROM files WHERE is_pdf = 1 AND is_corrupt = 0"
            " AND is_cloud_placeholder = 0 ORDER BY bates_filename_num, rel_path"
        )
        if args.limit:
            rows = rows[: args.limit]
        extraction_stats = _extract_documents(
            config, database, rows, run_id=run_id,
            progress=not args.no_progress, force=args.force,
        )
        _run_ocr_pass(config, database, [int(row["id"]) for row in rows], run_id=run_id)

        stats = _derive_and_report(config, database, rows, run_id=run_id, progress=not args.no_progress)
        stats.update(
            {
                "run_id": run_id,
                "pages_examined": extraction_stats.get("pages", 0),
                "extraction_failures": extraction_stats.get("failed", 0),
            }
        )
        database.finish_run(run_id, {k: v for k, v in stats.items() if isinstance(v, (str, int, float))})

    print(f"\nPhase 3 complete. Log: {log_path}")
    return 0


def _derive_and_report(
    config: Config,
    database: Database,
    rows: Sequence[sqlite3.Row],
    *,
    run_id: str,
    progress: bool = True,
) -> dict[str, Any]:
    """Derive metadata, group duplicates, infer families, and write reports."""
    from tqdm import tqdm

    classification_rules = config.get("classification.types", {}) or {}
    min_type_confidence = float(config.get("classification.min_confidence_to_assign", 0.25))
    compiled_tags = issue_tags.compile_terms(config.get("issue_tags.terms", {}) or {})
    max_hits = int(config.get("issue_tags.max_hits_per_tag", 25))
    context_chars = int(config.get("issue_tags.context_chars", 0))
    dup_config = config.section("duplicates")
    prefix = str(config.get("bates.prefix", "WILSONRODE"))
    digits = int(config.get("bates.expected_digits", 6))

    document_records: list[dict[str, Any]] = []
    issue_records: list[dict[str, Any]] = []
    family_candidates: list[document_family.FamilyCandidate] = []
    text_fingerprints: dict[int, str] = {}
    simhashes: dict[int, int] = {}
    order_key: dict[int, str] = {}
    text_lengths: dict[int, int] = {}
    embedded_counts: dict[int, int] = {}
    type_by_file: dict[int, str] = {}
    title_by_file: dict[int, str] = {}
    file_rows: dict[int, sqlite3.Row] = {int(row["id"]): row for row in rows}

    bar = tqdm(total=len(rows), disable=not progress, unit="doc", desc="Deriving metadata")

    for row in rows:
        file_id = int(row["id"])
        rel_path = str(row["rel_path"])
        order_key[file_id] = rel_path

        text = database.get_text(file_id) or ""
        page_rows = list(database.iter_pages(file_id))
        observed = [
            int(page["bates_observed_num"])
            for page in page_rows
            if page["bates_observed_num"] is not None
        ]
        conditions = [str(page["condition"]) for page in page_rows]
        searchable_pages = sum(1 for c in conditions if c == "Searchable")
        page_count = int(row["page_count"] or len(page_rows) or 0)
        is_blank = bool(page_rows) and all(
            condition in {"Apparent blank", "Bates-only page"} for condition in conditions
        )
        if page_count == 0:
            searchable_status = "Unknown"
        elif searchable_pages == page_count:
            searchable_status = "Fully searchable"
        elif searchable_pages == 0:
            searchable_status = "No searchable text"
        else:
            searchable_status = "Partially searchable"

        begin, end, source, confidence, reason = bates_mod.resolve_document_range(
            observed=observed,
            filename_number=row["bates_filename_num"],
            page_count=page_count or None,
        )

        metadata = metadata_extract.derive_metadata(
            text,
            classification_rules=classification_rules,
            min_type_confidence=min_type_confidence,
            is_blank=is_blank,
            searchable_status=searchable_status,
        )

        tag_result = issue_tags.tag_document(
            text,
            compiled_tags,
            max_hits_per_tag=max_hits,
            context_chars=context_chars,
        )
        for hit in tag_result.hits:
            issue_records.append(
                {
                    "file_id": file_id,
                    "tag": hit.tag,
                    "term": hit.term,
                    "page_number": hit.page_number,
                    "hit_count": hit.hit_count,
                    "locator": hit.locator,
                }
            )

        normalized = duplicates.normalize_text(
            text,
            lowercase=bool(dup_config.get("normalize_lowercase", True)),
            collapse_whitespace=bool(dup_config.get("normalize_collapse_whitespace", True)),
            strip_bates=bool(dup_config.get("normalize_strip_bates", True)),
            strip_digits=bool(dup_config.get("normalize_strip_digits", False)),
        )
        text_lengths[file_id] = len(normalized)
        if len(normalized) >= int(dup_config.get("min_chars_for_text_dupe", 200)):
            text_fingerprints[file_id] = duplicates.text_fingerprint(normalized)
            simhashes[file_id] = duplicates.simhash(
                normalized,
                bits=int(dup_config.get("simhash_bits", 64)),
                shingle_words=int(dup_config.get("shingle_words", 5)),
            )

        embedded_counts[file_id] = metadata.email.embedded_message_count
        type_by_file[file_id] = metadata.document_type
        title_by_file[file_id] = metadata.title or ""

        family_candidates.append(
            document_family.FamilyCandidate(
                file_id=file_id,
                rel_path=rel_path,
                filename=str(row["filename"]),
                begin_bates=begin,
                end_bates=end,
                document_type=metadata.document_type,
                title=metadata.title,
                subject=metadata.email.subject,
                document_date=metadata.document_date,
                attachment_names=tuple(metadata.attachment_names),
                page_count=page_count,
            )
        )

        document_records.append(
            {
                "file_id": file_id,
                "begin_bates": bates_mod.format_bates(begin, prefix=prefix, digits=digits),
                "end_bates": bates_mod.format_bates(end, prefix=prefix, digits=digits),
                "begin_bates_num": begin,
                "end_bates_num": end,
                "bates_source": source.value,
                "bates_confidence": confidence.value,
                "page_count": page_count,
                "document_date": metadata.document_date,
                "document_date_confidence": metadata.document_date_confidence,
                "email_sent_date": metadata.email.sent_date,
                "email_from": metadata.email.sender,
                "email_to": metadata.email.to_joined,
                "email_cc": metadata.email.cc_joined,
                "email_bcc": metadata.email.bcc_joined,
                "email_subject": metadata.email.subject,
                "embedded_message_count": metadata.email.embedded_message_count,
                "earliest_message_date": metadata.email.earliest_message_date,
                "latest_message_date": metadata.email.latest_message_date,
                "author": metadata.author,
                "title": metadata.title,
                "court": metadata.court,
                "case_number": metadata.case_number,
                "filing_date": metadata.filing_date,
                "filing_party": metadata.filing_party,
                "invoice_date": metadata.invoice_date,
                "vendor": metadata.vendor,
                "invoice_amount": metadata.invoice_amount,
                "document_type": metadata.document_type,
                "document_type_confidence": metadata.document_type_confidence,
                "attachment_names": "; ".join(metadata.attachment_names) or None,
                "custodian": metadata.custodian,
                "custodian_source": metadata.custodian_source,
                "searchable_status": searchable_status,
                "ocr_status": str(row["ocr_status"]),
                "apparent_blank_status": "Apparent blank or separator" if is_blank else "Has content",
                "issue_tags": tag_result.joined or None,
                "extraction_confidence": metadata.extraction_confidence,
                "manual_review": int(metadata.manual_review),
                "processing_notes": "; ".join([reason, *metadata.notes])[:2000],
                "app_version": APP_VERSION,
                "updated_ts": utcnow(),
            }
        )
        bar.update(1)
    bar.close()

    # -- duplicates ------------------------------------------------------
    LOG.info("Grouping duplicates", extra={"documents": len(document_records)})
    binary_members = duplicates.build_exact_binary_groups(
        [{"id": int(r["id"]), "sha256": r["sha256"], "rel_path": r["rel_path"]} for r in rows]
    )
    text_members = duplicates.build_text_groups(text_fingerprints, order_key=order_key)
    near_pairs = duplicates.find_near_duplicates(
        simhashes,
        bits=int(dup_config.get("simhash_bits", 64)),
        bands=int(dup_config.get("simhash_bands", 8)),
        max_distance=int(dup_config.get("simhash_max_distance", 3)),
        max_candidates_per_document=int(dup_config.get("max_candidates_per_document", 200)),
    )
    # Documents already grouped as exact or text duplicates do not need a
    # weaker near-duplicate label as well.
    exact_ids = {member.file_id for member in binary_members} | {
        member.file_id for member in text_members
    }
    near_pairs = [
        pair for pair in near_pairs if pair[0] not in exact_ids or pair[1] not in exact_ids
    ]
    near_members = duplicates.build_near_duplicate_groups(
        near_pairs, bits=int(dup_config.get("simhash_bits", 64)), order_key=order_key
    )

    all_members = [*binary_members, *text_members, *near_members]
    database.replace_rows(
        "duplicate_members",
        [
            {
                "group_id": member.group_id,
                "method": member.method,
                "file_id": member.file_id,
                "is_primary": int(member.is_primary),
                "relation": member.relation,
                "similarity": member.similarity,
            }
            for member in all_members
        ],
    )

    group_by_method: dict[str, dict[int, tuple[str, str, float | None, bool]]] = {
        duplicates.DuplicateMethod.EXACT_BINARY.value: {},
        duplicates.DuplicateMethod.NORMALIZED_TEXT.value: {},
        duplicates.DuplicateMethod.NEAR_DUPLICATE.value: {},
    }
    for member in all_members:
        group_by_method[member.method][member.file_id] = (
            member.group_id, member.relation, member.similarity, member.is_primary
        )

    # -- families --------------------------------------------------------
    relationships = document_family.infer_families(
        family_candidates,
        config=config.section("family"),
        bates_prefix=prefix,
        bates_digits=digits,
    )
    database.replace_rows(
        "families",
        [
            {
                "parent_file_id": rel.parent_file_id,
                "child_file_id": rel.child_file_id,
                "parent_bates": rel.parent_bates,
                "child_bates": rel.child_bates,
                "relationship_type": rel.relationship_type,
                "confidence": rel.confidence,
                "reason": rel.reason,
            }
            for rel in relationships
        ],
    )
    family_summary = document_family.summarise_for_document(relationships)

    # -- priority --------------------------------------------------------
    priority_config = config.section("priority")
    for record in document_records:
        file_id = int(record["file_id"])
        binary = group_by_method[duplicates.DuplicateMethod.EXACT_BINARY.value].get(file_id)
        text_group = group_by_method[duplicates.DuplicateMethod.NORMALIZED_TEXT.value].get(file_id)
        near_group = group_by_method[duplicates.DuplicateMethod.NEAR_DUPLICATE.value].get(file_id)

        record["exact_dup_group"] = binary[0] if binary else None
        record["text_dup_group"] = text_group[0] if text_group else None
        record["near_dup_group"] = near_group[0] if near_group else None

        summary = family_summary.get(file_id, {})
        record["_parent_bates"] = summary.get("parent_bates", "")
        record["_attachment_bates"] = summary.get("attachment_bates", "")
        record["_family_confidence"] = summary.get("family_confidence", "")

        result = issue_tags.score_document(
            tags=(record["issue_tags"] or "").split("; ") if record["issue_tags"] else [],
            config=priority_config,
            email_from=record["email_from"],
            email_to=record["email_to"],
            email_cc=record["email_cc"],
            document_date=record["document_date"],
            is_duplicate_secondary=bool(binary and not binary[3]),
            is_unique=not (binary or text_group or near_group),
            manual_review=bool(record["manual_review"]),
        )
        record["priority_score"] = result.score
        record["priority_explanation"] = result.explanation

    database.replace_rows("issue_hits", issue_records)
    database.replace_rows(
        "documents",
        [
            {key: value for key, value in record.items() if not key.startswith("_")}
            for record in document_records
        ],
    )

    # -- report rows -----------------------------------------------------
    master_rows: list[dict[str, Any]] = []
    for record in document_records:
        row = file_rows[int(record["file_id"])]
        master_rows.append(
            {
                "Begin Bates": record["begin_bates"],
                "End Bates": record["end_bates"],
                "Bates confidence": record["bates_confidence"],
                "Bates source": record["bates_source"],
                "Filename": row["filename"],
                "Relative path": row["rel_path"],
                "Open file": "Open",
                "Page count": record["page_count"],
                "File size (bytes)": row["file_size"],
                "File hash": row["sha256"],
                "Document date": record["document_date"],
                "Date confidence": record["document_date_confidence"],
                "Author or From": record["author"] or record["email_from"],
                "To": record["email_to"],
                "CC": record["email_cc"],
                "BCC": record["email_bcc"],
                "Subject or Title": record["email_subject"] or record["title"],
                "Document type": record["document_type"],
                "Document type confidence": record["document_type_confidence"],
                "Court": record["court"],
                "Case number": record["case_number"],
                "Parent Bates": record["_parent_bates"],
                "Attachment Bates": record["_attachment_bates"],
                "Family confidence": record["_family_confidence"],
                "Exact duplicate group": record["exact_dup_group"],
                "Text duplicate group": record["text_dup_group"],
                "Near duplicate group": record["near_dup_group"],
                "Issue tags": record["issue_tags"],
                "Priority score": record["priority_score"],
                "Priority score explanation": record["priority_explanation"],
                "Searchable status": record["searchable_status"],
                "OCR status": record["ocr_status"],
                "Apparent blank status": record["apparent_blank_status"],
                "Extraction confidence": record["extraction_confidence"],
                "Manual review required": "Yes" if record["manual_review"] else "No",
                "Processing notes": record["processing_notes"],
                "_abs_path": row["abs_path"],
            }
        )
    master_rows.sort(key=lambda item: (item["Begin Bates"] or "", item["Filename"]))

    ranges = [
        bates_mod.BatesRange(
            begin=int(record["begin_bates_num"]),
            end=int(record["end_bates_num"]),
            rel_path=str(file_rows[int(record["file_id"])]["rel_path"]),
        )
        for record in document_records
        if record["begin_bates_num"] is not None and record["end_bates_num"] is not None
    ]
    gaps, overlaps = bates_mod.find_range_gaps(ranges)
    gap_rows = [
        {
            "Gap after Bates": bates_mod.format_bates(gap.after, prefix=prefix, digits=digits),
            "Gap before Bates": bates_mod.format_bates(gap.before, prefix=prefix, digits=digits),
            "Missing numbers": gap.label,
            "Count missing": gap.missing_count,
            "Basis": gap.basis,
        }
        for gap in gaps
    ]
    overlap_rows = [
        {
            "Document A": left.rel_path,
            "A range": f"{left.begin}-{left.end}",
            "Document B": right.rel_path,
            "B range": f"{right.begin}-{right.end}",
            "Overlap": f"{max(left.begin, right.begin)}-{min(left.end, right.end)}",
            "Note": "Two documents claim the same Bates numbers; verify the page counts",
        }
        for left, right in overlaps
    ]

    duplicate_rows = []
    for member in all_members:
        row = file_rows.get(member.file_id)
        if row is None:
            continue
        record = next(
            (r for r in document_records if int(r["file_id"]) == member.file_id), None
        )
        relation = duplicates.refine_relation(
            method=duplicates.DuplicateMethod(member.method),
            left_chars=text_lengths.get(member.file_id, 0),
            right_chars=text_lengths.get(member.file_id, 0),
            left_embedded=embedded_counts.get(member.file_id, 0),
            right_embedded=0,
            left_type=type_by_file.get(member.file_id),
            right_type=None,
            left_title=title_by_file.get(member.file_id),
            right_title=None,
        ) if member.method == duplicates.DuplicateMethod.NEAR_DUPLICATE.value else None
        duplicate_rows.append(
            {
                "Group": member.group_id,
                "Method": member.method,
                "Relation": relation.value if relation else member.relation,
                "Similarity": member.similarity,
                "Primary": "Yes" if member.is_primary else "No",
                "Begin Bates": record["begin_bates"] if record else None,
                "Filename": row["filename"],
                "Relative path": row["rel_path"],
                "Open file": "Open",
                "Page count": row["page_count"],
                "_abs_path": row["abs_path"],
            }
        )
    duplicate_rows.sort(key=lambda item: (str(item["Method"]), str(item["Group"])))

    family_rows = [
        {
            "Parent Bates": rel.parent_bates,
            "Child Bates": rel.child_bates,
            "Relationship type": rel.relationship_type,
            "Confidence": rel.confidence,
            "Reason": rel.reason,
            "Parent filename": file_rows[rel.parent_file_id]["filename"]
            if rel.parent_file_id in file_rows else None,
            "Child filename": file_rows[rel.child_file_id]["filename"]
            if rel.child_file_id in file_rows else None,
            "Open parent": "Open",
            "_abs_path": file_rows[rel.parent_file_id]["abs_path"]
            if rel.parent_file_id in file_rows else None,
        }
        for rel in relationships
    ]

    bates_by_file = {
        int(record["file_id"]): record["begin_bates"] for record in document_records
    }
    issue_rows = [
        {
            "Issue tag": record["tag"],
            "Matching term": record["term"],
            "Hits": record["hit_count"],
            "Page": record["page_number"],
            "Locator": record["locator"],
            "Begin Bates": bates_by_file.get(int(record["file_id"])),
            "Filename": file_rows[int(record["file_id"])]["filename"],
            "Relative path": file_rows[int(record["file_id"])]["rel_path"],
            "Open file": "Open",
            "_abs_path": file_rows[int(record["file_id"])]["abs_path"],
        }
        for record in issue_records
        if int(record["file_id"]) in file_rows
    ]
    issue_rows.sort(key=lambda item: (str(item["Issue tag"]), str(item["Begin Bates"] or "")))

    blank_rows = []
    for record in document_records:
        row = file_rows[int(record["file_id"])]
        if record["apparent_blank_status"] != "Apparent blank or separator" and \
                record["searchable_status"] != "No searchable text":
            continue
        blank_rows.append(
            {
                "Begin Bates": record["begin_bates"],
                "Filename": row["filename"],
                "Relative path": row["rel_path"],
                "Open file": "Open",
                "Page count": record["page_count"],
                "Condition": record["apparent_blank_status"]
                if record["apparent_blank_status"] == "Apparent blank or separator"
                else "No searchable text",
                "Detail": record["processing_notes"],
                "Recommended action": "Visually confirm before treating this as blank",
                "_abs_path": row["abs_path"],
            }
        )

    error_rows = [
        {
            "Filename": row["filename"],
            "Relative path": row["rel_path"],
            "Open file": "Open",
            "Stage": "extraction" if row["extraction_status"] == STATUS_ERROR else "inventory",
            "Status": row["processing_status"],
            "Error": row["error_message"] or row["pdf_open_error"],
            "Recommended action": "Re-run with 'wri rerun-failed' once the cause is addressed",
            "_abs_path": row["abs_path"],
        }
        for row in database.iter_files(
            where="processing_status = 'error' OR extraction_status = 'error'"
            " OR is_corrupt = 1 OR is_cloud_placeholder = 1"
        )
    ]

    stats: dict[str, Any] = {
        "run_id": run_id,
        "documents_indexed": len(document_records),
        "documents_flagged_manual_review": sum(
            1 for record in document_records if record["manual_review"]
        ),
        "bates_gaps": len(gap_rows),
        "bates_overlaps": len(overlap_rows),
        "duplicate_members": len(all_members),
        "family_relationships": len(relationships),
        "issue_tag_hits": len(issue_records),
        "document_types": dict(Counter(r["document_type"] for r in document_records)),
        "bates_confidence": dict(Counter(r["bates_confidence"] for r in document_records)),
        "issue_tag_counts": dict(Counter(r["tag"] for r in issue_records)),
        "duplicate_counts": dict(Counter(m.method for m in all_members)),
        "family_confidence": dict(Counter(r.confidence for r in relationships)),
    }

    reports.write_phase3_reports(
        config,
        database,
        stats,
        master_rows=master_rows,
        gap_rows=gap_rows,
        overlap_rows=overlap_rows,
        duplicate_rows=duplicate_rows,
        family_rows=family_rows,
        issue_rows=issue_rows,
        blank_rows=blank_rows,
        error_rows=error_rows,
    )

    print("\n" + "=" * 72)
    print("PHASE 3 COMPLETE - DERIVED DOCUMENT METADATA")
    print("=" * 72)
    print(f"  Documents indexed ........... {stats['documents_indexed']:,}")
    print(f"  Flagged for manual review ... {stats['documents_flagged_manual_review']:,}")
    print(f"  Bates gaps (range-based) .... {stats['bates_gaps']:,}")
    print(f"  Bates overlaps .............. {stats['bates_overlaps']:,}")
    print(f"  Duplicate group memberships . {stats['duplicate_members']:,}")
    print(f"  Inferred family links ....... {stats['family_relationships']:,}")
    print(f"  Issue-tag hits .............. {stats['issue_tag_hits']:,}")
    print("\n  Document types:")
    for label, count in sorted(stats["document_types"].items(), key=lambda kv: -kv[1]):
        print(f"    {label:<26} {count:,}")
    print("=" * 72)
    return stats


# ---------------------------------------------------------------------------
# qc / search / status / rerun
# ---------------------------------------------------------------------------


def command_qc(args: argparse.Namespace) -> int:
    """Emit a deterministic quality-control sample for manual inspection."""
    config = _prepare(args, require_source=False)
    setup_logging(config.paths.logs_dir, run_name="qc")
    sample_config = config.get("qc.full_run_sample", {}) or {}
    seed = int(config.get("qc.random_seed", 20240101))

    with Database(config.paths.db_path) as database:
        buckets: list[tuple[str, str, int]] = [
            ("high priority", "SELECT d.file_id FROM documents d ORDER BY d.priority_score DESC",
             int(sample_config.get("high_priority", 50))),
            ("low priority", "SELECT d.file_id FROM documents d ORDER BY d.priority_score ASC",
             int(sample_config.get("low_priority", 50))),
            ("OCR", "SELECT f.id AS file_id FROM files f WHERE f.ocr_status IN ('completed','required','failed')",
             int(sample_config.get("ocr", 25))),
            ("blank candidates",
             "SELECT DISTINCT p.file_id FROM pages p WHERE p.is_blank_candidate = 1",
             int(sample_config.get("blank_candidates", 25))),
            ("family inferences", "SELECT child_file_id AS file_id FROM families",
             int(sample_config.get("family_inferences", 25))),
            ("near duplicates",
             "SELECT file_id FROM duplicate_members WHERE method = 'near_duplicate'",
             int(sample_config.get("near_duplicates", 25))),
        ]

        rng = random.Random(seed)
        rows: list[dict[str, Any]] = []
        for label, sql, quota in buckets:
            candidates = [int(record["file_id"]) for record in database.query(sql)]
            if len(candidates) > quota:
                candidates = rng.sample(candidates, quota) if "ORDER BY" not in sql else candidates[:quota]
            for file_id in candidates:
                record = database.query(
                    "SELECT f.filename, f.rel_path, f.abs_path, d.begin_bates, d.document_type,"
                    " d.priority_score, d.extraction_confidence, d.manual_review"
                    " FROM files f LEFT JOIN documents d ON d.file_id = f.id WHERE f.id = ?",
                    (file_id,),
                )
                if not record:
                    continue
                item = record[0]
                rows.append(
                    {
                        "QC bucket": label,
                        "Begin Bates": item["begin_bates"],
                        "Filename": item["filename"],
                        "Relative path": item["rel_path"],
                        "Open file": "Open",
                        "Document type": item["document_type"],
                        "Priority score": item["priority_score"],
                        "Extraction confidence": item["extraction_confidence"],
                        "Manual review": "Yes" if item["manual_review"] else "No",
                        "Reviewer verdict": "",
                        "Reviewer notes": "",
                        "_abs_path": item["abs_path"],
                    }
                )

        from openpyxl import Workbook

        workbook = Workbook()
        columns = [
            "QC bucket", "Begin Bates", "Filename", "Relative path", "Open file",
            "Document type", "Priority score", "Extraction confidence", "Manual review",
            "Reviewer verdict", "Reviewer notes",
        ]
        reports.write_sheet(
            workbook, "QC Sample", columns, rows, config=config, hyperlink_column="Open file"
        )
        path = reports.finalise_workbook(
            workbook,
            config.paths.reports_dir / "QC_Sample.xlsx",
            specs=(
                reports.ColumnSpec("QC bucket", reports.CALCULATED, "Which stratum this row was drawn from.", 22),
                reports.ColumnSpec("Reviewer verdict", reports.OBSERVED, "Fill in during manual inspection: correct / incorrect / unclear.", 20),
                reports.ColumnSpec("Reviewer notes", reports.OBSERVED, "Free-text reviewer observations.", 50, wrap=True),
            ),
            stats={"sample_rows": len(rows), "seed": seed},
            extra_notes=[
                f"This sample is deterministic for seed {seed}: re-running produces "
                "the same rows, so a QC pass can be reproduced and audited.",
            ],
        )

    print(f"QC sample written: {path}")
    print(f"Rows: {len(rows)}")
    print("Open each linked document and record a verdict. Do not treat any derived")
    print("field as verified until it has been checked against the document itself.")
    return 0


def command_search(args: argparse.Namespace) -> int:
    """Search the FTS5 index and print file identities only."""
    config = _prepare(args, require_source=False)
    with Database(config.paths.db_path, read_only=True) as database:
        try:
            results = database.search_text(args.query, limit=args.limit)
        except sqlite3.OperationalError as exc:
            print(f"Search failed: {exc}")
            return 1

    if not results:
        print("No matches.")
        return 0
    print(f"{len(results)} match(es). Document text is deliberately not displayed.\n")
    for row in results:
        print(f"  {row['rel_path']}")
    return 0


def command_status(args: argparse.Namespace) -> int:
    """Print processing progress across all phases."""
    config = _prepare(args, require_source=False)
    with Database(config.paths.db_path, read_only=True) as database:
        total = database.count_files()
        if total == 0:
            print("No inventory yet. Run 'wri phase1'.")
            return 0
        print(f"Database: {config.paths.db_path}")
        print(f"  files inventoried .......... {total:,}")
        print(f"  inventory complete ......... {database.count_files('processing_status = ?', (STATUS_DONE,)):,}")
        print(f"  inventory errors ........... {database.count_files('processing_status = ?', (STATUS_ERROR,)):,}")
        print(f"  extraction complete ........ {database.count_files('extraction_status = ?', (STATUS_DONE,)):,}")
        print(f"  extraction errors .......... {database.count_files('extraction_status = ?', (STATUS_ERROR,)):,}")
        print(f"  pages examined ............. {database.scalar('SELECT COUNT(*) FROM pages', default=0):,}")
        print(f"  documents with metadata .... {database.scalar('SELECT COUNT(*) FROM documents', default=0):,}")
        print(f"  cloud placeholders ......... {database.count_files('is_cloud_placeholder = 1'):,}")
        for run in database.query("SELECT * FROM runs ORDER BY started_ts DESC LIMIT 5"):
            print(f"\n  run {run['run_id']} ({run['phase']}) {run['status']}")
            print(f"    started  {run['started_ts']}")
            print(f"    finished {run['finished_ts'] or '-'}")
    return 0


def command_rerun_failed(args: argparse.Namespace) -> int:
    """Reset failed files so the next run reprocesses only those."""
    config = _prepare(args, require_source=False)
    setup_logging(config.paths.logs_dir, run_name="rerun")
    with Database(config.paths.db_path) as database:
        rows = database.query(
            "SELECT id, rel_path FROM files WHERE processing_status = 'error'"
            " OR extraction_status = 'error' OR is_cloud_placeholder = 1"
        )
        if not rows:
            print("No failed files to requeue.")
            return 0
        file_ids = [int(row["id"]) for row in rows]
        database.delete_file_derivatives(file_ids)
        with database.transaction() as conn:
            marks = ",".join("?" for _ in file_ids)
            conn.execute(
                f"UPDATE files SET processing_status='pending', extraction_status='pending',"
                f" error_message=NULL, seen_size=NULL, seen_mtime=NULL WHERE id IN ({marks})",
                file_ids,
            )
        database.audit("rerun_failed_requeued", detail={"count": len(file_ids)})
    print(f"Requeued {len(rows)} file(s). Re-run 'wri phase1' (and later phases) to process them.")
    return 0


# ---------------------------------------------------------------------------
# shared plumbing
# ---------------------------------------------------------------------------


def _prepare(args: argparse.Namespace, *, require_source: bool = True) -> Config:
    """Load configuration, apply CLI overrides, and enforce path safety."""
    overrides: dict[str, Any] = {}
    if getattr(args, "source", None):
        overrides["paths.source_folder"] = args.source
    if getattr(args, "output", None):
        overrides["paths.output_folder"] = args.output
    if getattr(args, "workers", None):
        overrides["performance.workers"] = args.workers

    config = load_config(args.config, overrides=overrides)
    config.validate_paths(require_source=require_source)
    config.paths.ensure_output_tree()
    return config


def build_parser() -> argparse.ArgumentParser:
    """Construct the argument parser for every command."""
    parser = argparse.ArgumentParser(
        prog="wri",
        description=(
            f"{APP_NAME} v{APP_VERSION} - local, read-only indexing and triage for a "
            "Bates-numbered document production. All output is DERIVED review "
            "metadata, never native metadata."
        ),
    )
    parser.add_argument("--config", help="Path to config.yaml", default=None)
    parser.add_argument("--source", help="Override paths.source_folder", default=None)
    parser.add_argument("--output", help="Override paths.output_folder", default=None)
    parser.add_argument("--workers", type=int, help="Override the worker count", default=None)
    parser.add_argument("--no-progress", action="store_true", help="Disable progress bars")
    parser.add_argument("--version", action="version", version=f"{APP_NAME} {APP_VERSION}")

    subparsers = parser.add_subparsers(dest="command", required=True)

    locate = subparsers.add_parser("locate", help="Find the source production folder")
    locate.add_argument("--name", default=TARGET_FOLDER_NAME, help="Exact folder name to find")
    locate.set_defaults(func=command_locate)

    doctor = subparsers.add_parser("doctor", help="Check dependencies and configuration")
    doctor.set_defaults(func=command_doctor)

    phase1 = subparsers.add_parser("phase1", help="Structural inventory (no text is read)")
    phase1.add_argument("--force", action="store_true", help="Re-probe already-processed files")
    phase1.add_argument("--limit", type=int, default=None, help="Stop after N files")
    phase1.set_defaults(func=command_phase1)

    phase2 = subparsers.add_parser("phase2", help="Text and page-condition audit")
    phase2.add_argument("--full", action="store_true", help="Process the full production")
    phase2.add_argument("--sample", type=int, default=None, help="Pilot sample size")
    phase2.add_argument("--force", action="store_true", help="Re-extract completed files")
    phase2.set_defaults(func=command_phase2)

    phase3 = subparsers.add_parser("phase3", help="Derived metadata and review workbooks")
    phase3.add_argument("--force", action="store_true", help="Re-extract completed files")
    phase3.add_argument("--limit", type=int, default=None, help="Process only N documents")
    phase3.set_defaults(func=command_phase3)

    qc = subparsers.add_parser("qc", help="Write the quality-control sample workbook")
    qc.set_defaults(func=command_qc)

    search = subparsers.add_parser("search", help="Search the FTS5 index (identities only)")
    search.add_argument("query", help="FTS5 MATCH expression")
    search.add_argument("--limit", type=int, default=50)
    search.set_defaults(func=command_search)

    status = subparsers.add_parser("status", help="Show processing progress")
    status.set_defaults(func=command_status)

    rerun = subparsers.add_parser("rerun-failed", help="Requeue only the failed files")
    rerun.set_defaults(func=command_rerun_failed)

    return parser


def main(argv: Sequence[str] | None = None) -> int:
    """Entry point. Returns a process exit code."""
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        return int(args.func(args))
    except ConfigError as exc:
        print(f"\nConfiguration problem:\n  {exc}\n", file=sys.stderr)
        return 2
    except KeyboardInterrupt:
        print("\nInterrupted. Progress is committed; re-run the same command to resume.")
        return 130


if __name__ == "__main__":
    raise SystemExit(main())
